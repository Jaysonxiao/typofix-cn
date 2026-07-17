from docx import Document
import pytest

from typofix_cn.documents.docx_reader import DocxReader


def test_reader_extracts_body_and_table_but_not_header(tmp_path) -> None:
    path = tmp_path / "论文.docx"
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    paragraph = document.add_paragraph("这是正文。")
    paragraph.style = document.styles["Normal"]
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格文本。"
    document.sections[0].header.paragraphs[0].text = "页眉不检查"
    document.save(path)

    blocks = DocxReader().read(path, relative_path="论文.docx")

    assert [block.text for block in blocks] == ["第一章 绪论", "这是正文。", "表格文本。"]
    assert blocks[0].role == "heading"
    assert blocks[1].role == "body"
    assert blocks[2].region == "table_cell"
    assert blocks[2].table.row_index == 0


def test_reader_concatenates_runs_and_classifies_lists_and_blank_paragraphs(tmp_path) -> None:
    path = tmp_path / "runs.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("跨")
    paragraph.add_run("Run")
    document.add_paragraph("")
    list_paragraph = document.add_paragraph("列表项")
    list_paragraph.style = "List Bullet"
    document.save(path)

    blocks = DocxReader().read(path, relative_path="runs.docx")

    assert blocks[0].text == "跨Run"
    assert blocks[1].role == "blank"
    assert blocks[2].role == "list_item"


def test_reader_rejects_non_docx_zip(tmp_path) -> None:
    path = tmp_path / "not-docx.docx"
    path.write_bytes(b"PK\x03\x04not a Word document")

    with pytest.raises(Exception, match="DOCX"):
        DocxReader().read(path, relative_path="not-docx.docx")
